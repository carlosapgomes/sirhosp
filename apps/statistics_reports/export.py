"""In-memory XLSX export of one daily statistics revision (DSRS-S7).

The exporter reads exactly the projection the authorized page renders, so the
workbook can never disagree with the screen it mirrors:

- every rendered group becomes one worksheet, in the order the page renders it,
  and the conditional report-level group becomes the ``Setor não identificado``
  worksheet without pretending to be an official grouping;
- every worksheet keeps the fixed sections in the required order -- admissions,
  transfer arrivals, deaths, transfer departures, hospital discharges, events
  with an unidentified endpoint and the closing patients -- including the empty
  ones, each showing its own count beside the title;
- rows reuse the wording, the fields and the natural ordering the projection
  already computed for the page, and a missing value stays missing instead of
  being filled in;
- a value that starts with a formula-significant character is stored as text,
  and a sheet label that Excel would reject or collide with is normalized,
  truncated and de-duplicated deterministically while the full grouping name
  stays inside the sheet;
- the workbook is built in memory and handed over as bytes: no path is opened,
  no file is persisted and the download name carries only the date and the
  revision.

This module is read-only with respect to the database: it never writes, never
materializes and never audits. The served-export audit row is committed by the
view only after the workbook exists.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from apps.statistics_reports.models import (
    DailyStatisticsPatient,
    DailyStatisticsReport,
)
from apps.statistics_reports.presentation import (
    PATIENT_LIST_TITLE,
    UNIDENTIFIED_LIST_TITLE,
    DailyReportProjection,
    EventRow,
    ReportGroup,
)

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
"""Media type of the workbook served by the export endpoint."""

SHEET_TITLE_LIMIT = 31
"""Maximum worksheet name length Excel accepts."""

FALLBACK_SHEET_TITLE = "Setor"
"""Worksheet name used when a grouping label keeps no usable character."""

UNKNOWN_SHEET_MARKER = "sem agrupamento oficial"
"""Cell that identifies the conditional worksheet of a non-official group."""

FILENAME_PREFIX = "estatisticas-diarias"
"""Filename stem; the served name carries only the date and the revision."""

EVENT_COLUMNS: tuple[str, ...] = (
    "Leito",
    "Nome",
    "Prontuário",
    "Tipo",
    "Momento clínico",
    "Origem",
    "Destino",
    "Detecção",
)
"""Event columns, in the same meaning and order the page renders them."""

PATIENT_COLUMNS: tuple[str, ...] = (
    "Leito",
    "Nome",
    "Prontuário",
    "Especialidade",
)
"""Closing patient columns, in the same meaning and order the page renders."""

_INVALID_SHEET_TITLE_CHARS = frozenset("[]:*?/\\")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_BOLD = Font(bold=True)


@dataclass(frozen=True)
class SheetSection:
    """One fixed section every worksheet carries, empty or not."""

    title: str
    attribute: str
    columns: tuple[str, ...]
    holds_patients: bool = False


SECTIONS: tuple[SheetSection, ...] = (
    SheetSection("Internações", "admissions", EVENT_COLUMNS),
    SheetSection("Transferências de entrada", "transfer_entries", EVENT_COLUMNS),
    SheetSection("Óbitos", "deaths", EVENT_COLUMNS),
    SheetSection("Transferências de saída", "transfer_exits", EVENT_COLUMNS),
    SheetSection("Altas hospitalares", "discharges", EVENT_COLUMNS),
    SheetSection(UNIDENTIFIED_LIST_TITLE, "unidentified", EVENT_COLUMNS),
    SheetSection(PATIENT_LIST_TITLE, "patients", PATIENT_COLUMNS, holds_patients=True),
)
"""Fixed sections of every worksheet, in the order the report requires."""


@dataclass(frozen=True)
class ExportWorkbook:
    """One workbook built in memory and ready to be served."""

    filename: str
    content: bytes
    sheet_count: int
    row_count: int


def build_export_workbook(
    projection: DailyReportProjection,
) -> ExportWorkbook:
    """Build the workbook of one rendered revision, entirely in memory.

    Every rendered group becomes one worksheet in the order the page renders
    it; the conditional report-level group becomes the ``Setor não identificado``
    worksheet that mirrors the page section. No filesystem path is ever used
    and the returned bytes are the whole, finalized workbook.
    """
    workbook = Workbook()
    titles = sheet_titles([group.title for group in projection.groups])
    worksheets: list[Worksheet] = []
    for title in titles:
        if worksheets:
            worksheets.append(workbook.create_sheet(title=title))
        else:
            first = workbook.active
            first.title = title
            worksheets.append(first)

    row_count = 0
    for worksheet, group in zip(worksheets, projection.groups, strict=True):
        row_count += _write_group_sheet(
            worksheet=worksheet, group=group, report=projection.report
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return ExportWorkbook(
        filename=export_filename(projection.report),
        content=buffer.getvalue(),
        sheet_count=len(workbook.sheetnames),
        row_count=row_count,
    )


def export_filename(report: DailyStatisticsReport) -> str:
    """Safe download name built only from the date and the revision."""
    return (
        f"{FILENAME_PREFIX}-{report.local_date.isoformat()}-r{report.revision}.xlsx"
    )


def sheet_titles(titles: Sequence[str]) -> tuple[str, ...]:
    """Valid, unique and deterministic worksheet names for the given labels.

    A label is normalized to what Excel accepts, truncated to the worksheet
    limit, and a collision -- including one the normalization itself created --
    is resolved with a stable numeric suffix while the full name stays inside
    the sheet.
    """
    used: set[str] = set()
    resolved: list[str] = []
    for title in titles:
        base = _sanitize_sheet_title(title)
        candidate = base
        index = 2
        while candidate.casefold() in used:
            suffix = f" ({index})"
            candidate = f"{base[: SHEET_TITLE_LIMIT - len(suffix)]}{suffix}"
            index += 1
        used.add(candidate.casefold())
        resolved.append(candidate)
    return tuple(resolved)


def _sanitize_sheet_title(title: str) -> str:
    """One label reduced to what Excel accepts as a worksheet name."""
    cleaned = "".join(
        " " if character in _INVALID_SHEET_TITLE_CHARS else character
        for character in title
    )
    cleaned = cleaned.strip().strip("'").strip()
    if not cleaned:
        return FALLBACK_SHEET_TITLE
    return cleaned[:SHEET_TITLE_LIMIT]


def _write_group_sheet(
    *, worksheet: Worksheet, group: ReportGroup, report: DailyStatisticsReport
) -> int:
    """Write one grouping worksheet and return its aggregate data row count."""
    _write_cell(worksheet, 1, 1, group.title, bold=True)
    _write_cell(
        worksheet,
        1,
        2,
        group.stable_key or UNKNOWN_SHEET_MARKER,
        bold=True,
    )
    _write_cell(worksheet, 2, 1, f"Data do relatório: {report.local_date.isoformat()}")
    _write_cell(worksheet, 2, 2, f"Revisão: {report.revision}")

    cursor = 4
    data_rows = 0
    for section in SECTIONS:
        entries = getattr(group, section.attribute)
        _write_cell(worksheet, cursor, 1, section.title, bold=True)
        _write_cell(worksheet, cursor, 2, len(entries), bold=True)
        cursor += 1
        for column, header in enumerate(section.columns, start=1):
            _write_cell(worksheet, cursor, column, header, bold=True)
        cursor += 1
        for entry in entries:
            values = (
                _patient_values(entry)
                if section.holds_patients
                else _event_values(entry)
            )
            for column, value in enumerate(values, start=1):
                _write_cell(worksheet, cursor, column, value)
            cursor += 1
            data_rows += 1
        cursor += 1
    return data_rows


def _write_cell(
    worksheet: Worksheet, row: int, column: int, value: object, *, bold: bool = False
) -> None:
    """Write one cell, keeping text starting a formula inert.

    The workbook must never evaluate nominal content: a text value whose first
    character is formula-significant is stored with the explicit string type
    instead of the formula type openpyxl would otherwise assign.
    """
    cell = worksheet.cell(row=row, column=column, value=value)
    if bold:
        cell.font = _BOLD
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        cell.data_type = "s"


def _event_values(row: EventRow) -> tuple[str, ...]:
    """Nominal and temporal fields of one event row, as the page shows them."""
    event = row.event
    return (
        event.bed,
        event.name,
        event.record,
        row.label,
        row.clinical_display,
        row.origin_display,
        row.destination_display,
        row.detection_display,
    )


def _patient_values(patient: DailyStatisticsPatient) -> tuple[str, ...]:
    """Nominal fields of one closing patient row, as the page shows them."""
    return (patient.bed, patient.name, patient.record, patient.specialty)
