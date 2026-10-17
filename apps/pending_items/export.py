"""XLSX workbook for the aggregate cut (PEND-S4, D8/PP-08).

Pure serialization: the view selects rows through the shared S3 queries
and this module formats them into an in-memory workbook. Every value is
a literal string cell (``data_type == 's'``), so formula-looking text
(``=``, ``+``, ``-``, ``@``) can never execute. Dates use the
institutional timezone; system-driven updates always read ``Sistema``.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import TYPE_CHECKING, Sequence

import openpyxl
from django.utils import timezone
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.pending_items.services import SYSTEM_AUTHOR_NAME

if TYPE_CHECKING:
    from apps.pending_items.models import PendingItem

EXPORT_HEADERS: tuple[str, ...] = (
    "Registro",
    "Nome",
    "Unidade atual",
    "Leito",
    "Primeira unidade de internação",
    "Especialidade",
    "Tipo",
    "Descrição/observação",
    "Status",
    "Criado em",
    "Criado por",
    "Atualizado em",
    "Atualizado por",
)
"""Exact PP-08 column order; also fixes the expected row width."""

EXPORT_SHEET_TITLE = "Pendências"

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

EXPORT_FILENAME_PREFIX = "pendencias"

EXPORT_DATETIME_FORMAT = "%d/%m/%Y %H:%M"
"""Same wall-clock rendering the aggregate panel uses (``d/m/Y H:i``)."""

UNKNOWN_AUTHOR_LABEL = "—"
"""Fallback mirroring the aggregate panel for missing author names."""


def format_export_datetime(value: datetime) -> str:
    """Render an instant in the institutional timezone for Excel."""
    return timezone.localtime(value).strftime(EXPORT_DATETIME_FORMAT)


def created_by_label(item: PendingItem) -> str:
    """Human creation author, never blank in the workbook."""
    return item.created_by_name or UNKNOWN_AUTHOR_LABEL


def updated_by_label(item: PendingItem) -> str:
    """Last-update author; automatic updates read ``Sistema`` always."""
    if item.is_system_updated:
        return SYSTEM_AUTHOR_NAME
    return item.updated_by_name or UNKNOWN_AUTHOR_LABEL


def build_pending_items_workbook(rows: Sequence[Sequence[str]]) -> bytes:
    """Serialize header plus one row per item to XLSX bytes.

    All values are forced to literal string cells; the workbook lives
    only in memory (``BytesIO``) with no temporary file on disk.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = EXPORT_SHEET_TITLE
    header_font = Font(bold=True)
    for col_idx, header in enumerate(EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.data_type = "s"
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = sheet.cell(
                row=row_idx, column=col_idx, value=value or ""
            )
            cell.data_type = "s"
    for col_idx, header in enumerate(EXPORT_HEADERS, start=1):
        width = len(header)
        for row in sheet.iter_rows(
            min_row=2, min_col=col_idx, max_col=col_idx
        ):
            for cell in row:
                if cell.value:
                    width = max(width, len(str(cell.value)))
        sheet.column_dimensions[get_column_letter(col_idx)].width = (
            width + 3
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_filename(stamp: datetime) -> str:
    """Censo-style filename stamp (``pendencias-YYYYMMDD-HHMM.xlsx``)."""
    file_ts = timezone.localtime(stamp).strftime("%Y%m%d-%H%M")
    return f"{EXPORT_FILENAME_PREFIX}-{file_ts}.xlsx"
