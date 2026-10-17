"""PEND-S4: pending-items XLSX workbook unit tests (R3/R4 workbook core).

Failing-before: ``apps/pending_items/export.py`` does not exist yet, so
importing this module fails until the slice is implemented. Pure checks
(no database): exact PP-08 column order, literal text cells by TYPE
(``data_type`` in ``s``/``inlineStr``, never formula) including
formula-looking payloads, empty
workbook validity, institutional-timezone dates and author labels with
``Sistema`` for automatic updates. All fixtures are synthetic.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from django.template.loader import get_template
from django.urls import resolve, reverse

from apps.pending_items.export import (
    EXPORT_HEADERS,
    XLSX_MIME,
    build_pending_items_workbook,
    created_by_label,
    format_export_datetime,
    updated_by_label,
)
from apps.pending_items.models import PendingItem

TZ = ZoneInfo("America/Bahia")

EXPECTED_HEADERS = [
    "Registro",
    "Nome",
    "Unidade atual",
    "Leito",
    "Primeira unidade de internação",
    "Especialidade",
    "Tipo",
    "Descrição/observação",
    "Status",
    "Criado em",
    "Criado por",
    "Atualizado em",
    "Atualizado por",
]

FORMULA_LOOKALIKES = ["=1+1", "+1+1", "-2+3", "@SOMA(A1:A2)"]


def _reopen(payload: bytes):
    return openpyxl.load_workbook(BytesIO(payload)).active


def _assert_literal(cell) -> None:
    assert cell.data_type in ("s", "inlineStr")


class TestExportHeaders:
    def test_headers_match_pp08_exact_order(self) -> None:
        assert list(EXPORT_HEADERS) == EXPECTED_HEADERS

    def test_mime_constant(self) -> None:
        assert XLSX_MIME == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    def test_empty_workbook_has_headers_only(self) -> None:
        ws = _reopen(build_pending_items_workbook([]))
        assert ws is not None
        assert [cell.value for cell in ws[1]] == EXPECTED_HEADERS
        assert ws.max_row == 1


class TestLiteralTextCells:
    def test_every_text_value_is_literal_string_type(self) -> None:
        row = [f"valor-{index}" for index in range(len(EXPECTED_HEADERS))]
        ws = _reopen(build_pending_items_workbook([row]))
        assert ws is not None
        assert ws.max_row == 2
        for cell in ws[2]:
            _assert_literal(cell)
        assert [cell.value for cell in ws[2]] == row

    def test_formula_lookalikes_stay_literal_in_every_column(self) -> None:
        base = FORMULA_LOOKALIKES * 4
        row = base[: len(EXPECTED_HEADERS)]
        ws = _reopen(build_pending_items_workbook([row]))
        assert ws is not None
        for cell in ws[2]:
            _assert_literal(cell)
        assert [cell.value for cell in ws[2]] == row

    def test_empty_values_round_trip_as_blank_literal_cells(self) -> None:
        row = ["" for _ in range(len(EXPECTED_HEADERS))]
        ws = _reopen(build_pending_items_workbook([row]))
        assert ws is not None
        assert ws.max_row == 2
        for cell in ws[2]:
            _assert_literal(cell)
            assert cell.value in ("", None)

    def test_one_row_per_record_without_concatenation(self) -> None:
        rows = [
            [f"primeira-{index}" for index in range(len(EXPECTED_HEADERS))],
            [f"segunda-{index}" for index in range(len(EXPECTED_HEADERS))],
        ]
        ws = _reopen(build_pending_items_workbook(rows))
        assert ws is not None
        assert ws.max_row == 3
        assert [cell.value for cell in ws[2]] == rows[0]
        assert [cell.value for cell in ws[3]] == rows[1]


class TestExportFormatting:
    def test_datetime_uses_institutional_timezone(self) -> None:
        utc = datetime(2026, 6, 10, 15, 0, 0, tzinfo=ZoneInfo("UTC"))
        assert format_export_datetime(utc) == "10/06/2026 12:00"

    def test_system_update_never_shows_human_name(self) -> None:
        item = PendingItem(
            is_system_updated=True, updated_by_name="humano.sintetico"
        )
        assert updated_by_label(item) == "Sistema"

    def test_human_update_shows_recorded_name(self) -> None:
        item = PendingItem(
            is_system_updated=False, updated_by_name="humano.sintetico"
        )
        assert updated_by_label(item) == "humano.sintetico"

    def test_missing_authors_fall_back_to_panel_placeholder(self) -> None:
        item = PendingItem(
            is_system_updated=False,
            created_by_name="",
            updated_by_name="",
        )
        assert created_by_label(item) == "—"
        assert updated_by_label(item) == "—"


class TestExportWiring:
    def test_route_name_resolves_to_xlsx_path(self) -> None:
        assert reverse("pending_items:aggregate_export") == (
            "/pendencias/export.xlsx"
        )

    def test_path_dispatches_to_authenticated_export_view(self) -> None:
        match = resolve("/pendencias/export.xlsx")
        assert match.func.__name__ == "aggregate_export"

    def test_aggregate_panel_links_export_with_filters(self) -> None:
        template = get_template("pending_items/aggregate.html")
        origin = getattr(template, "origin", None)
        assert origin is not None
        source = Path(origin.name).read_text(encoding="utf-8")
        assert "pending_items:aggregate_export" in source
        assert 'data-testid="aggregate-export"' in source
        assert "request.GET.urlencode" in source
