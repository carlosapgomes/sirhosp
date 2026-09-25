"""DSRS-S7 integration tests: XLSX export and served-export audit.

Covers the vertical slice requirements:

- R1: the export endpoint requires ``export_daily_statistics`` independently of
  the consultation permission and serves a ``private, no-store`` response;
- R2: the workbook has one worksheet per official grouping, valid and unique
  sheet names with the full grouping name inside the sheet, plus the
  conditional ``Setor não identificado`` worksheet whenever the revision
  carries events with neither endpoint or closing patients without an
  attributable sector, neither of which is ever omitted;
- R3: every sheet keeps the fixed sections in order -- admissions, transfer
  arrivals, deaths, transfer departures, hospital discharges, events with an
  unidentified endpoint and the closing patients;
- R4: empty sections stay present with a zero count and rows keep the fields
  and the natural ordering of the page;
- R5: text that starts with a formula-significant character is stored as safe
  text;
- R6: the workbook is built in memory, no file is persisted and the filename
  depends only on the date and the revision;
- R7: a served workbook records user, instant, selected date, revision and
  aggregate counts, while a failure before the prepared response records
  nothing;
- R8: the audit log carries no nominal payload.

Everything here uses synthetic runs, sectors, beds and patients; no real
extraction data, no production access and no persisted workbook.
"""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest
from django.contrib.auth.models import Permission, User
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from apps.census.models import CapacityCatalogVersion
from apps.statistics_reports.export import (
    XLSX_CONTENT_TYPE,
    build_export_workbook,
)
from apps.statistics_reports.models import DailyStatisticsReport, StatisticsExportLog
from apps.statistics_reports.presentation import (
    DailyReportProjection,
    ReportGroup,
    build_daily_report_projection,
)
from tests.integration.test_daily_statistics_materialization import (
    _base_lines,
    _build_catalog,
    _closing_lines,
)
from tests.integration.test_daily_statistics_page import (
    DAY,
    DEATH_NAME,
    DEATH_PATIENT,
    EXPORT_PERMISSION_CODENAME,
    GERAL_CODE,
    GERAL_KEY,
    GERAL_SECTOR,
    MOVING_NAME,
    NO_GROUPING_CODE,
    NO_GROUPING_NAME,
    NO_GROUPING_PATIENT,
    UNATTRIBUTED_NAME,
    UNATTRIBUTED_PATIENT,
    UNKNOWN_SECTION_TITLE,
    VIEW_PERMISSION_CODENAME,
    WIDE_DAY,
    _anchor_lines,
    _closing_lines_with_ordering,
    _death_evidence,
    _death_lines,
    _group_of,
    _materialize_day,
    _patient_line,
    _wide_day_payload,
    page_url,
)

EXPORT_URL_NAME = "statistics_reports:daily_report_export"

EXPECTED_SECTIONS = (
    "Internações",
    "Transferências de entrada",
    "Óbitos",
    "Transferências de saída",
    "Altas hospitalares",
    "Eventos com origem ou destino não identificado",
    "Pacientes do fechamento",
)

EVENT_HEADERS = [
    "Leito",
    "Nome",
    "Prontuário",
    "Tipo",
    "Momento clínico",
    "Origem",
    "Destino",
    "Detecção",
]

PATIENT_HEADERS = ["Leito", "Nome", "Prontuário", "Especialidade"]

GERAL_SHEET = "Gerais"


def export_url() -> str:
    return reverse(EXPORT_URL_NAME)


def _workbook(response: Any) -> Workbook:
    return load_workbook(BytesIO(response.content))


def _section(worksheet: Worksheet, title: str) -> tuple[int, list[object], list[list[object]]]:
    """Count cell, header row and data rows of one fixed section of a sheet."""
    for row_index in range(1, worksheet.max_row + 1):
        if worksheet.cell(row=row_index, column=1).value != title:
            continue
        count = worksheet.cell(row=row_index, column=2).value
        header = [
            worksheet.cell(row=row_index + 1, column=column).value
            for column in range(1, worksheet.max_column + 1)
        ]
        width = len([value for value in header if value is not None])
        rows: list[list[object]] = []
        cursor = row_index + 2
        while cursor <= worksheet.max_row:
            first = worksheet.cell(row=cursor, column=1).value
            if first is None or first in EXPECTED_SECTIONS:
                break
            rows.append(
                [worksheet.cell(row=cursor, column=column).value for column in range(1, width + 1)]
            )
            cursor += 1
        return int(count), header[:width], rows
    raise AssertionError(f"section {title!r} not found in {worksheet.title!r}")


def _sheet_sections(worksheet: Worksheet) -> list[str]:
    """Every section title of one sheet, in the order it is written."""
    return [
        worksheet.cell(row=row, column=1).value
        for row in range(1, worksheet.max_row + 1)
        if worksheet.cell(row=row, column=1).value in EXPECTED_SECTIONS
    ]


def _data_row_count(workbook: Workbook) -> int:
    """Aggregate data rows of one workbook, section titles excluded."""
    total = 0
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        for title in EXPECTED_SECTIONS:
            total += len(_section(worksheet, title)[2])
    return total


NO_GROUPING_SECTOR = "SETOR FORA DO CATALOGO"
NO_GROUPING_BED = "999-A"


def _patient_only_unattributable_report(
    catalog: CapacityCatalogVersion,
) -> DailyStatisticsReport:
    """A day whose only unattributable row is one ungrouped closing patient.

    The both-endpoint-null anchor patient is removed and the ungrouped patient
    is present in the opening and the closing photograph, so the revision has
    no event whose endpoints are both null and its sole unattributable row is
    that closing patient.
    """
    ungrouped = _patient_line(
        code=NO_GROUPING_CODE,
        sector=NO_GROUPING_SECTOR,
        bed=NO_GROUPING_BED,
        record=NO_GROUPING_PATIENT,
        name=NO_GROUPING_NAME,
    )
    opening = [
        line for line in _anchor_lines() if line.record != UNATTRIBUTED_PATIENT
    ]
    opening.append(ungrouped)
    return _materialize_day(
        catalog=catalog,
        local_date=DAY,
        lines=opening,
        closing_lines=_closing_lines_with_ordering(without_grouping=True),
    )


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def catalog(db: None) -> CapacityCatalogVersion:
    """Synthetic historical catalog of the shared DSRS-S2 fixtures."""
    return _build_catalog(
        effective_from=date(2026, 1, 1),
        algorithm_version="occupancy-v5",
        source_reference="synthetic DSRS-S7 catalog",
        source_sha256="e" * 64,
    )


@pytest.fixture
def report(catalog: CapacityCatalogVersion) -> DailyStatisticsReport:
    """The rich synthetic day: movements, unknown endpoints and ordering beds."""
    return _materialize_day(
        catalog=catalog,
        local_date=DAY,
        lines=_anchor_lines(),
        closing_lines=_closing_lines_with_ordering(),
    )


def _permission(codename: str) -> Permission:
    return Permission.objects.get(
        codename=codename,
        content_type__app_label="statistics_reports",
    )


def _logged_in_client(*, username: str, codenames: tuple[str, ...]) -> Client:
    user = User.objects.create_user(username=username, password="testpass123")
    user.user_permissions.add(*[_permission(name) for name in codenames])
    client = Client()
    client.login(username=username, password="testpass123")
    return client


@pytest.fixture
def export_client(db: None) -> Client:
    """A client holding both the consultation and the export permission."""
    return _logged_in_client(
        username="exportador",
        codenames=(VIEW_PERMISSION_CODENAME, EXPORT_PERMISSION_CODENAME),
    )


@pytest.fixture
def export_only_client(db: None) -> Client:
    """A client holding only the export permission, without consultation."""
    return _logged_in_client(
        username="somente-exportacao",
        codenames=(EXPORT_PERMISSION_CODENAME,),
    )


@pytest.fixture
def view_only_client(db: None) -> Client:
    """A client holding only the consultation permission."""
    return _logged_in_client(
        username="somente-consulta",
        codenames=(VIEW_PERMISSION_CODENAME,),
    )


@pytest.fixture
def plain_client(db: None) -> Client:
    """A separate client logged in as a user WITHOUT any report permission."""
    return _logged_in_client(username="comum", codenames=())


# ---------------------------------------------------------------------------
# R1 - independent authorization and sensitive response policy
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestExportAuthorization:
    def test_anonymous_request_redirects_to_login(
        self, client: Client, report: DailyStatisticsReport
    ) -> None:
        response = client.get(export_url())
        assert response.status_code == 302
        assert reverse("login") in response["Location"]

    def test_view_permission_alone_is_not_enough(
        self, view_only_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = view_only_client.get(export_url(), {"date": DAY.isoformat()})
        assert response.status_code == 403
        assert StatisticsExportLog.objects.count() == 0

    def test_user_without_any_permission_is_denied(
        self, plain_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = plain_client.get(export_url(), {"date": DAY.isoformat()})
        assert response.status_code == 403
        assert StatisticsExportLog.objects.count() == 0

    def test_export_permission_alone_allows_the_download(
        self, export_only_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = export_only_client.get(export_url(), {"date": DAY.isoformat()})
        assert response.status_code == 200
        assert response["Content-Type"] == XLSX_CONTENT_TYPE

    def test_served_workbook_is_private_and_not_stored(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = export_client.get(export_url(), {"date": DAY.isoformat()})
        assert response.status_code == 200
        assert response["Cache-Control"] == "private, no-store"

    def test_unavailable_date_serves_no_workbook_and_logs_nothing(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = export_client.get(export_url(), {"date": (DAY + timedelta(days=2)).isoformat()})
        assert response.status_code == 404
        assert StatisticsExportLog.objects.count() == 0

    def test_page_offers_the_export_link_with_the_export_permission(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = export_client.get(page_url(), {"date": DAY.isoformat()})
        body = response.content.decode()
        assert response.status_code == 200
        assert f'href="{export_url()}?date={DAY.isoformat()}"' in body
        assert "Exportar XLSX" in body

    def test_page_hides_the_export_link_without_the_export_permission(
        self, view_only_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = view_only_client.get(page_url(), {"date": DAY.isoformat()})
        body = response.content.decode()
        assert response.status_code == 200
        assert "Exportar XLSX" not in body
        assert export_url() not in body


# ---------------------------------------------------------------------------
# R6 - in-memory workbook and safe filename
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestWorkbookDelivery:
    def test_workbook_is_built_in_memory_only(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        saved: list[Any] = []
        real_save = Workbook.save

        def spy(*args: Any, **kwargs: Any) -> None:
            target = args[1] if len(args) > 1 else kwargs.get("filename")
            saved.append(target)
            real_save(*args, **kwargs)

        with patch.object(Workbook, "save", new=spy):
            response = export_client.get(export_url(), {"date": DAY.isoformat()})
        assert response.status_code == 200
        assert saved
        assert all(isinstance(target, BytesIO) for target in saved)

    def test_filename_depends_only_on_date_and_revision(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = export_client.get(export_url(), {"date": DAY.isoformat()})
        expected = (
            f'attachment; filename="estatisticas-diarias-{DAY.isoformat()}-r{report.revision}.xlsx"'
        )
        assert response["Content-Disposition"] == expected
        assert MOVING_NAME not in response["Content-Disposition"]

    def test_no_export_file_is_left_behind(
        self, export_client: Client, report: DailyStatisticsReport, tmp_path: Path
    ) -> None:
        response = export_client.get(export_url(), {"date": DAY.isoformat()})
        filename = response["Content-Disposition"].split('filename="')[1][:-1]
        assert response.status_code == 200
        assert filename.endswith(".xlsx")
        assert not (tmp_path / filename).exists()
        assert not (Path.cwd() / filename).exists()


# ---------------------------------------------------------------------------
# R2 - one sheet per grouping plus the conditional unknown sheet
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestWorkbookSheets:
    def test_one_sheet_per_rendered_grouping(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        projection = build_daily_report_projection(report)
        response = export_client.get(export_url(), {"date": DAY.isoformat()})
        workbook = _workbook(response)
        assert workbook.sheetnames == [group.title for group in projection.groups]

    def test_sheet_identifies_its_full_grouping_name(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        projection = build_daily_report_projection(report)
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        for group in projection.groups:
            worksheet = workbook[group.title]
            assert worksheet.cell(row=1, column=1).value == group.title
            expected_key = group.stable_key or "sem agrupamento oficial"
            assert worksheet.cell(row=1, column=2).value == expected_key

    def test_sheet_names_are_valid_unique_and_keep_the_full_name(
        self, report: DailyStatisticsReport
    ) -> None:
        long_title = "SETOR " + "MUITO LONGO " * 4
        projection = DailyReportProjection(
            report=report,
            groups=(
                ReportGroup(
                    key="a",
                    stable_key="A",
                    title="SETOR COLIDE:1",
                    is_unknown_section=False,
                    header_badges=(),
                ),
                ReportGroup(
                    key="b",
                    stable_key="B",
                    title="SETOR COLIDE/1",
                    is_unknown_section=False,
                    header_badges=(),
                ),
                ReportGroup(
                    key="c",
                    stable_key="C",
                    title=long_title,
                    is_unknown_section=False,
                    header_badges=(),
                ),
                ReportGroup(
                    key="d",
                    stable_key="D",
                    title="SETOR [INVALIDO]?",
                    is_unknown_section=False,
                    header_badges=(),
                ),
            ),
        )
        workbook = load_workbook(BytesIO(build_export_workbook(projection).content))
        names = workbook.sheetnames
        assert len(names) == 4
        assert len(set(names)) == 4
        for name in names:
            assert len(name) <= 31
            assert not set(name) & set("[]:*?/\\")
        assert names[0] == "SETOR COLIDE 1"
        assert names[1] == "SETOR COLIDE 1 (2)"
        assert workbook[names[2]].cell(row=1, column=1).value == long_title

    def test_unknown_section_becomes_a_conditional_sheet(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        assert UNKNOWN_SECTION_TITLE in workbook.sheetnames
        unknown = workbook[UNKNOWN_SECTION_TITLE]
        values = [cell.value for row in unknown.iter_rows() for cell in row]
        assert UNATTRIBUTED_NAME in values
        for name in workbook.sheetnames:
            if name == UNKNOWN_SECTION_TITLE:
                continue
            other = workbook[name]
            assert UNATTRIBUTED_NAME not in [
                cell.value for row in other.iter_rows() for cell in row
            ]

    def test_unattributable_patient_alone_creates_the_conditional_sheet(
        self, export_client: Client, catalog: CapacityCatalogVersion
    ) -> None:
        report = _patient_only_unattributable_report(catalog)
        assert not report.events.filter(
            origin_sector__isnull=True, destination_sector__isnull=True
        ).exists()
        projection = build_daily_report_projection(report)
        assert projection.unknown_section is not None
        assert projection.unknown_section.unidentified == ()
        assert [
            patient.record for patient in projection.unknown_section.patients
        ] == [NO_GROUPING_PATIENT]

        workbook = _workbook(
            export_client.get(export_url(), {"date": DAY.isoformat()})
        )
        assert workbook.sheetnames == [group.title for group in projection.groups]
        assert workbook.sheetnames.count(UNKNOWN_SECTION_TITLE) == 1
        unknown = workbook[UNKNOWN_SECTION_TITLE]
        values = [cell.value for row in unknown.iter_rows() for cell in row]
        assert NO_GROUPING_NAME in values
        for name in workbook.sheetnames:
            if name == UNKNOWN_SECTION_TITLE:
                continue
            other = workbook[name]
            assert NO_GROUPING_NAME not in [
                cell.value for row in other.iter_rows() for cell in row
            ]

    def test_unknown_sheet_is_absent_without_unattributable_rows(
        self, catalog: CapacityCatalogVersion
    ) -> None:
        report = _materialize_day(
            catalog=catalog,
            local_date=DAY,
            lines=_base_lines(),
            closing_lines=_base_lines() + _closing_lines(),
        )
        assert not report.events.filter(
            origin_sector__isnull=True, destination_sector__isnull=True
        ).exists()
        assert not report.patients.filter(sector__isnull=True).exists()
        projection = build_daily_report_projection(report)
        assert projection.unknown_section is None
        workbook = load_workbook(BytesIO(build_export_workbook(projection).content))
        assert UNKNOWN_SECTION_TITLE not in workbook.sheetnames

    def test_event_without_any_endpoint_is_never_omitted(
        self, export_client: Client, catalog: CapacityCatalogVersion
    ) -> None:
        _death_evidence(record=DEATH_PATIENT, name=DEATH_NAME)
        _materialize_day(
            catalog=catalog,
            local_date=DAY,
            lines=_death_lines(),
            closing_lines=_closing_lines_with_ordering(),
        )
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        unknown = workbook[UNKNOWN_SECTION_TITLE]
        values = [cell.value for row in unknown.iter_rows() for cell in row]
        assert DEATH_NAME in values
        assert DEATH_PATIENT in values


# ---------------------------------------------------------------------------
# R3/R4 - fixed sections, counts, fields and natural ordering
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestSheetSections:
    def test_every_sheet_keeps_the_fixed_sections_in_order(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        assert workbook.sheetnames
        for name in workbook.sheetnames:
            assert tuple(_sheet_sections(workbook[name])) == EXPECTED_SECTIONS

    def test_empty_sections_stay_present_with_a_zero_count(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        worksheet = workbook[GERAL_SHEET]
        count, headers, rows = _section(worksheet, "Internações")
        assert count == 0
        assert rows == []
        assert headers == EVENT_HEADERS

    def test_section_counts_match_the_page_projection(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        projection = build_daily_report_projection(report)
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        for group in projection.groups:
            worksheet = workbook[group.title]
            assert _section(worksheet, "Internações")[0] == len(group.admissions)
            assert _section(worksheet, "Transferências de entrada")[0] == len(
                group.transfer_entries
            )
            assert _section(worksheet, "Óbitos")[0] == len(group.deaths)
            assert _section(worksheet, "Transferências de saída")[0] == len(group.transfer_exits)
            assert _section(worksheet, "Altas hospitalares")[0] == len(group.discharges)
            assert _section(worksheet, "Eventos com origem ou destino não identificado")[0] == len(
                group.unidentified
            )
            assert _section(worksheet, "Pacientes do fechamento")[0] == len(group.patients)

    def test_closing_patients_keep_the_page_fields_and_order(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        projection = build_daily_report_projection(report)
        group = _group_of(projection, GERAL_KEY)
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        _, headers, rows = _section(workbook[group.title], "Pacientes do fechamento")
        assert headers == PATIENT_HEADERS
        assert [row[1] for row in rows] == [patient.name for patient in group.patients]
        assert [row[0] for row in rows] == [patient.bed for patient in group.patients]
        assert [row[2] for row in rows] == [patient.record for patient in group.patients]
        assert [row[3] for row in rows] == [patient.specialty for patient in group.patients]

    def test_event_rows_keep_the_page_fields_and_order(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        projection = build_daily_report_projection(report)
        group = _group_of(projection, GERAL_KEY)
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        _, headers, rows = _section(
            workbook[group.title],
            "Eventos com origem ou destino não identificado",
        )
        assert headers == EVENT_HEADERS
        assert rows
        assert [row[0] for row in rows] == [row.event.bed for row in group.unidentified]
        assert [row[1] for row in rows] == [row.event.name for row in group.unidentified]
        assert [row[2] for row in rows] == [row.event.record for row in group.unidentified]
        assert [row[3] for row in rows] == [rendered.label for rendered in group.unidentified]
        assert [row[4] for row in rows] == [
            rendered.clinical_display for rendered in group.unidentified
        ]
        assert [row[5] for row in rows] == [
            rendered.origin_display for rendered in group.unidentified
        ]
        assert [row[6] for row in rows] == [
            rendered.destination_display for rendered in group.unidentified
        ]
        assert [row[7] for row in rows] == [
            rendered.detection_display for rendered in group.unidentified
        ]


# ---------------------------------------------------------------------------
# R5 - formula-significant text is written as safe text
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestFormulaSafety:
    def test_formula_like_text_is_stored_as_text(
        self, export_client: Client, catalog: CapacityCatalogVersion
    ) -> None:
        dangerous = ("=1+1", "+1+1", "-1+1", "@SUM(A1)")
        extra = [
            _patient_line(
                code=GERAL_CODE,
                sector=GERAL_SECTOR,
                bed=f"D{index}",
                record=f"770{index}",
                name=name,
            )
            for index, name in enumerate(dangerous)
        ]
        _materialize_day(
            catalog=catalog,
            local_date=DAY,
            lines=_anchor_lines() + extra,
            closing_lines=_closing_lines_with_ordering() + extra,
        )
        workbook = _workbook(export_client.get(export_url(), {"date": DAY.isoformat()}))
        found: dict[str, str] = {}
        for row in workbook[GERAL_SHEET].iter_rows():
            for cell in row:
                if cell.value in dangerous:
                    found[cell.value] = cell.data_type
        assert set(found) == set(dangerous)
        assert all(data_type == "s" for data_type in found.values())


# ---------------------------------------------------------------------------
# R7/R8 - served-export audit log
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestExportAudit:
    def test_success_logs_user_date_revision_and_aggregate_counts(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        response = export_client.get(export_url(), {"date": DAY.isoformat()})
        workbook = _workbook(response)
        logs = list(StatisticsExportLog.objects.all())
        assert len(logs) == 1
        log = logs[0]
        assert log.user == User.objects.get(username="exportador")
        assert log.report == report
        assert log.report.local_date == DAY
        assert log.report.revision == report.revision
        assert log.sheet_count == len(workbook.sheetnames)
        assert log.row_count == _data_row_count(workbook)
        assert log.row_count > 0
        assert log.served_at is not None

    def test_failure_before_the_response_logs_no_success(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        with patch(
            "apps.statistics_reports.views.build_export_workbook",
            side_effect=RuntimeError("workbook generation failed"),
        ):
            with pytest.raises(RuntimeError):
                export_client.get(export_url(), {"date": DAY.isoformat()})
        assert StatisticsExportLog.objects.count() == 0

    def test_failure_constructing_the_response_logs_no_success(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        with patch(
            "apps.statistics_reports.views.HttpResponse",
            side_effect=RuntimeError("response construction failed"),
        ):
            with pytest.raises(RuntimeError):
                export_client.get(export_url(), {"date": DAY.isoformat()})
        assert StatisticsExportLog.objects.count() == 0

    def test_failure_configuring_the_response_headers_logs_no_success(
        self, export_client: Client, report: DailyStatisticsReport
    ) -> None:
        with patch("apps.statistics_reports.views.HttpResponse") as response_cls:
            response_cls.return_value.__setitem__.side_effect = RuntimeError(
                "response header configuration failed"
            )
            with pytest.raises(RuntimeError):
                export_client.get(export_url(), {"date": DAY.isoformat()})
        assert StatisticsExportLog.objects.count() == 0

    def test_audit_model_carries_no_nominal_payload(self) -> None:
        field_names = {field.name for field in StatisticsExportLog._meta.concrete_fields}
        assert field_names == {
            "id",
            "user",
            "report",
            "served_at",
            "sheet_count",
            "row_count",
        }


# ---------------------------------------------------------------------------
# Bounded read of the same projection as the page
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestExportQueryBudget:
    def test_query_count_does_not_grow_with_the_revision(
        self, export_client: Client, catalog: CapacityCatalogVersion, report: DailyStatisticsReport
    ) -> None:
        wide_lines, wide_closing, payload = _wide_day_payload()
        _materialize_day(
            catalog=catalog,
            local_date=WIDE_DAY,
            lines=wide_lines,
            closing_lines=wide_closing,
            groups=payload["groups"],
        )
        with CaptureQueriesContext(connection) as small_ctx:
            small = export_client.get(export_url(), {"date": DAY.isoformat()})
        with CaptureQueriesContext(connection) as big_ctx:
            big = export_client.get(export_url(), {"date": WIDE_DAY.isoformat()})
        assert small.status_code == 200
        assert big.status_code == 200
        assert len(_workbook(big).sheetnames) > len(_workbook(small).sheetnames)
        assert len(big_ctx) == len(small_ctx)
