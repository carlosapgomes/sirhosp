"""PEND-S4: pending-items XLSX export HTTP integration (R1-R4).

Failing-before: no ``pendencias/export.xlsx`` route exists in
``apps/pending_items/urls.py``, so every export GET below 404s until
the slice is implemented.

Covers authenticated-only download for readers, members and groupless
superusers, the shared conservative cut with one row per item across
pages and views, exact PP-08 columns with first unit/authors/dates and
no history, literal text cells, empty workbooks, cache directives, GET
purity and the censo-style filename. All synthetic.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import openpyxl
import pytest
from django.contrib.auth.models import Group, User
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from apps.census.models import BedStatus, CensusSnapshot, Specialty
from apps.patients.models import Admission, Patient
from apps.pending_items.export import EXPORT_HEADERS
from apps.pending_items.models import (
    OUTRO_SLUG,
    STATUS_ACTIVE,
    STATUS_CLOSED_UNRESOLVED,
    STATUS_RESOLVED,
    PendingItem,
    PendingItemRevision,
    PendingItemType,
)
from apps.pending_items.queries import AGGREGATE_PER_PAGE

TZ = ZoneInfo("America/Bahia")
EXPORT_URL = "/pendencias/export.xlsx"
EXPECTED_HEADERS = [
    "Registro",
    "Nome",
    "Unidade atual",
    "Leito",
    "Primeira unidade de internação",
    "Especialidade",
    "Tipo",
    "Descrição/observação",
    "Status",
    "Criado em",
    "Criado por",
    "Atualizado em",
    "Atualizado por",
]


def _dt(value: str) -> datetime:
    return datetime.fromisoformat(value).replace(tzinfo=TZ)


def _make_patient(key: str) -> Patient:
    return Patient.objects.create(
        patient_source_key=key,
        source_system="tasy",
        name=f"PACIENTE SINTETICO {key}",
    )


def _make_admission(
    patient: Patient,
    key: str,
    start: str | None = "2026-06-01T08:00:00",
    end: str | None = None,
) -> Admission:
    return Admission.objects.create(
        patient=patient,
        source_system="tasy",
        source_admission_key=key,
        admission_date=_dt(start) if start else None,
        discharge_date=_dt(end) if end else None,
    )


def _make_census_row(
    prontuario: str,
    *,
    setor: str = "CLINICA MEDICA",
    leito: str = "L01",
    especialidade: str = "",
    captured_at: str = "2026-06-10T08:00:00",
) -> CensusSnapshot:
    return CensusSnapshot.objects.create(
        captured_at=_dt(captured_at),
        setor=setor,
        leito=leito,
        prontuario=prontuario,
        nome="PACIENTE SINTETICO",
        especialidade=especialidade,
        bed_status=BedStatus.OCCUPIED,
    )


def _make_item(
    admission: Admission,
    *,
    description: str = "pendencia sintetica",
    status: str = STATUS_ACTIVE,
    created_at: datetime | None = None,
    finalized_at: datetime | None = None,
) -> PendingItem:
    outro = PendingItemType.objects.get(slug=OUTRO_SLUG)
    item = PendingItem.objects.create(
        admission=admission,
        item_type=outro,
        description=description,
        status=status,
        version=1,
        finalized_at=finalized_at,
    )
    if created_at is not None:
        PendingItem.objects.filter(pk=item.pk).update(created_at=created_at)
        item.refresh_from_db()
    return item


def _eligible_admission(
    key: str = "P_EXP_1",
    *,
    setor: str = "CLINICA MEDICA",
    leito: str = "L01",
    especialidade: str = "",
) -> Admission:
    patient = _make_patient(key)
    admission = _make_admission(patient, f"ADM_{key}")
    _make_census_row(
        key, setor=setor, leito=leito, especialidade=especialidade
    )
    return admission


def _workbook(response):
    assert response.status_code == 200
    return openpyxl.load_workbook(BytesIO(response.content)).active


def _rows(ws) -> list[list]:
    assert ws is not None
    return [
        [cell.value for cell in row] for row in ws.iter_rows(min_row=2)
    ]


def _descriptions(ws) -> list[str]:
    return [row[7] for row in _rows(ws)]


def _cache_directives(response) -> set[str]:
    header = response.headers.get("Cache-Control", "")
    return {part.strip().lower() for part in header.split(",") if part.strip()}


def _assert_private_no_store(response) -> None:
    directives = _cache_directives(response)
    assert "private" in directives
    assert "no-store" in directives


@pytest.fixture
def writer_client(client: Client, db: object) -> Client:
    user = User.objects.create_user(username="export_writer", password="pw123456")
    Group.objects.get(name="Desospitalização").user_set.add(user)
    client.login(username="export_writer", password="pw123456")
    return client


@pytest.fixture
def reader_client(client: Client, db: object) -> Client:
    User.objects.create_user(username="export_reader", password="pw123456")
    client.login(username="export_reader", password="pw123456")
    return client


@pytest.fixture
def groupless_superuser_client(client: Client, db: object) -> Client:
    User.objects.create_superuser(
        username="export_root", password="pw123456",
        email="root@example.com",
    )
    client.login(username="export_root", password="pw123456")
    return client


@pytest.mark.django_db
class TestExportAccess:
    def test_route_name_resolves_to_xlsx_path(self) -> None:
        assert reverse("pending_items:aggregate_export") == EXPORT_URL

    def test_reader_member_and_groupless_superuser_export(
        self,
        reader_client: Client,
        writer_client: Client,
        groupless_superuser_client: Client,
    ) -> None:
        admission = _eligible_admission("P_EXP_R1")
        _make_item(admission, description="exportavel por todos")
        for http_client in (
            reader_client, writer_client, groupless_superuser_client
        ):
            ws = _workbook(http_client.get(EXPORT_URL))
            assert "exportavel por todos" in _descriptions(ws)

    def test_export_grants_no_write_permission(
        self, reader_client: Client, groupless_superuser_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_R1W")
        item = _make_item(admission, description="protegida export")
        outro = PendingItemType.objects.get(slug=OUTRO_SLUG)
        for http_client in (reader_client, groupless_superuser_client):
            assert http_client.get(EXPORT_URL).status_code == 200
            response = http_client.post(
                f"/pendencias/{item.pk}/edit/",
                {
                    "item_type": str(outro.pk),
                    "description": "tentativa",
                    "expected_version": "1",
                    "next": "/pendencias/",
                },
            )
            assert response.status_code == 403
        item.refresh_from_db()
        assert item.description == "protegida export"

    def test_anonymous_gets_login_without_bytes(self, client: Client) -> None:
        admission = _eligible_admission("P_EXP_ANON")
        _make_item(admission, description="anonima export")
        response = client.get(EXPORT_URL)
        assert response.status_code == 302
        assert "/login/" in response.url  # type: ignore[attr-defined]
        assert b"anonima export" not in response.content


@pytest.mark.django_db
class TestExportCutParity:
    def test_all_results_not_the_visible_page(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_PAGE")
        for index in range(AGGREGATE_PER_PAGE + 1):
            _make_item(admission, description=f"pagina-{index:02d}")
        page1 = reader_client.get("/pendencias/?view=items&page=1")
        assert f"pagina-{AGGREGATE_PER_PAGE:02d}" not in (
            page1.content.decode()
        )
        ws = _workbook(reader_client.get(f"{EXPORT_URL}?view=items"))
        assert _descriptions(ws) == [
            f"pagina-{index:02d}" for index in range(AGGREGATE_PER_PAGE + 1)
        ]

    def test_mode_is_ignored_views_export_same_selection(
        self, reader_client: Client
    ) -> None:
        first = _eligible_admission(
            "P_EXP_MODE1", setor="UTI", leito="U01"
        )
        second = _eligible_admission(
            "P_EXP_MODE2", setor="ENF", leito="E02"
        )
        _make_item(first, description="modo primeira")
        _make_item(second, description="modo segunda")
        patients = _rows(
            _workbook(reader_client.get(f"{EXPORT_URL}?view=patients"))
        )
        items = _rows(
            _workbook(reader_client.get(f"{EXPORT_URL}?view=items"))
        )
        assert patients == items
        assert [row[7] for row in patients] == [
            "modo segunda",
            "modo primeira",
        ]

    def test_status_and_unit_filters_apply_before_export(
        self, reader_client: Client
    ) -> None:
        first = _eligible_admission("P_EXP_F1", setor="UTI", leito="U01")
        second = _eligible_admission("P_EXP_F2", setor="UTI", leito="U02")
        _make_item(first, description="leito um")
        _make_item(
            second, description="leito dois resolvido",
            status=STATUS_RESOLVED,
            finalized_at=_dt("2026-06-11T10:00:00"),
        )
        ws = _workbook(
            reader_client.get(f"{EXPORT_URL}?status=all&unidade=UTI&leito=U02")
        )
        assert _descriptions(ws) == ["leito dois resolvido"]
        ws = _workbook(reader_client.get(EXPORT_URL))
        assert _descriptions(ws) == ["leito um"]

    def test_system_updater_and_time_bounds_match_panel(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_F3")
        outro = PendingItemType.objects.get(slug=OUTRO_SLUG)
        now = timezone.now()
        PendingItem.objects.create(
            admission=admission,
            item_type=outro,
            description="encerrada automatica export",
            status=STATUS_CLOSED_UNRESOLVED,
            version=2,
            finalized_at=_dt("2026-06-11T10:00:00"),
            updated_by=None,
            updated_by_name="Sistema",
            is_system_updated=True,
        )
        _make_item(
            admission, description="antiga humana",
            created_at=now - timedelta(days=9),
        )
        _make_item(
            admission, description="recente humana",
            created_at=now - timedelta(days=1),
        )
        ws = _workbook(
            reader_client.get(f"{EXPORT_URL}?status=all&atualizador=system")
        )
        assert _descriptions(ws) == ["encerrada automatica export"]
        ws = _workbook(reader_client.get(f"{EXPORT_URL}?min_dias=5"))
        assert _descriptions(ws) == ["antiga humana"]
        ws = _workbook(
            reader_client.get(f"{EXPORT_URL}?max_dias=2&ordenar=tempo_desc")
        )
        assert _descriptions(ws) == ["recente humana"]

    def test_ambiguous_excluded_and_repetition_single(
        self, reader_client: Client
    ) -> None:
        conflict = _make_patient("P_EXP_AMB1")
        adm_conflict = _make_admission(conflict, "ADM_EXP_AMB1")
        _make_census_row("P_EXP_AMB1", setor="UTI", leito="U01")
        _make_census_row("P_EXP_AMB1", setor="ENF", leito="E02")
        _make_item(adm_conflict, description="leitos conflitantes export")
        two = _make_patient("P_EXP_AMB2")
        adm_two = _make_admission(two, "ADM_EXP_AMB2_A")
        _make_admission(two, "ADM_EXP_AMB2_B")
        _make_census_row("P_EXP_AMB2")
        _make_item(adm_two, description="duas abertas export")
        valid = _eligible_admission("P_EXP_AMB3")
        _make_census_row("P_EXP_AMB3")
        _make_item(valid, description="valida export")
        html = reader_client.get("/pendencias/?status=all").content.decode()
        ws = _workbook(reader_client.get(f"{EXPORT_URL}?status=all"))
        found = _descriptions(ws)
        assert "leitos conflitantes export" not in html
        assert "duas abertas export" not in html
        assert "leitos conflitantes export" not in found
        assert "duas abertas export" not in found
        assert found.count("valida export") == 1

    def test_invalid_time_bounds_export_headers_only(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_F4")
        _make_item(admission, description="visivel sem filtro export")
        ws = _workbook(reader_client.get(f"{EXPORT_URL}?min_dias=-2"))
        assert ws is not None
        assert ws.max_row == 1

    def test_ordering_matches_items_selection(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_ORD")
        now = timezone.now()
        _make_item(
            admission, description="ordem antiga",
            created_at=now - timedelta(days=9),
        )
        _make_item(
            admission, description="ordem recente",
            created_at=now - timedelta(days=1),
        )
        asc = _descriptions(
            _workbook(reader_client.get(f"{EXPORT_URL}?ordenar=tempo_asc"))
        )
        desc = _descriptions(
            _workbook(reader_client.get(f"{EXPORT_URL}?ordenar=tempo_desc"))
        )
        assert asc == ["ordem recente", "ordem antiga"]
        assert desc == ["ordem antiga", "ordem recente"]


@pytest.mark.django_db
class TestExportColumnsAndValues:
    def test_columns_values_authors_and_dates(
        self, reader_client: Client
    ) -> None:
        Specialty.objects.get_or_create(
            code="CIV", defaults={"name": "CIRURGIA VASCULAR"}
        )
        author = User.objects.get(username="export_reader")
        other = User.objects.create_user(
            username="export_editor", password="pw123456"
        )
        admission = _eligible_admission(
            "P_EXP_COL", setor="UTI GERAL", leito="U07",
            especialidade="CIV",
        )
        outro = PendingItemType.objects.get(slug=OUTRO_SLUG)
        item = PendingItem.objects.create(
            admission=admission,
            item_type=outro,
            description="linha completa export",
            status=STATUS_ACTIVE,
            version=1,
            created_by=author,
            created_by_name="export_reader",
            updated_by=other,
            updated_by_name="export_editor",
            is_system_updated=False,
        )
        PendingItem.objects.filter(pk=item.pk).update(
            created_at=_dt("2026-06-05T10:30:00"),
            updated_at=_dt("2026-06-09T18:45:00"),
        )
        ws = _workbook(reader_client.get(EXPORT_URL))
        assert ws is not None
        assert [cell.value for cell in ws[1]] == EXPECTED_HEADERS
        assert list(EXPORT_HEADERS) == EXPECTED_HEADERS
        rows = _rows(ws)
        assert len(rows) == 1
        assert rows[0] == [
            "P_EXP_COL",
            "PACIENTE SINTETICO P_EXP_COL",
            "UTI GERAL",
            "U07",
            "Primeira unidade não identificada",
            "CIRURGIA VASCULAR",
            "Outro",
            "linha completa export",
            "Ativa",
            "05/06/2026 10:30",
            "export_reader",
            "09/06/2026 18:45",
            "export_editor",
        ]

    def test_system_author_and_no_history_leak(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_SYS")
        outro = PendingItemType.objects.get(slug=OUTRO_SLUG)
        item = PendingItem.objects.create(
            admission=admission,
            item_type=outro,
            description="atual sistema export",
            status=STATUS_CLOSED_UNRESOLVED,
            version=2,
            finalized_at=_dt("2026-06-11T10:00:00"),
            updated_by=None,
            updated_by_name="Sistema",
            is_system_updated=True,
        )
        PendingItemRevision.objects.create(
            item=item, version=1, action="created", item_type=outro,
            type_name=outro.name, description="atual sistema export",
            status=STATUS_ACTIVE,
        )
        PendingItemRevision.objects.create(
            item=item, version=2, action="finalized", item_type=outro,
            type_name=outro.name, description="atual sistema export",
            status=STATUS_CLOSED_UNRESOLVED, is_system=True,
            author_name="Sistema", note="nota-somente-revisao-export",
        )
        ws = _workbook(reader_client.get(f"{EXPORT_URL}?status=all"))
        rows = _rows(ws)
        assert len(rows) == 1
        assert rows[0][8] == "Encerrada sem resolução"
        assert rows[0][12] == "Sistema"
        flat = " ".join(str(value) for row in rows for value in row)
        assert "nota-somente-revisao-export" not in flat

    def test_first_unit_uses_episode_pathway(
        self, reader_client: Client
    ) -> None:
        from apps.census.models import PatientMovement

        admission = _eligible_admission("P_EXP_FU")
        day = datetime.strptime("2026-06-02", "%Y-%m-%d").date()
        PatientMovement.objects.create(
            patient=admission.patient, admission=admission,
            movement_date=day, sector="PRONTO SOCORRO", sequence=0,
            first_seen_at=_dt("2026-06-01T08:00:00"),
            last_seen_at=_dt("2026-06-01T08:00:00"),
        )
        _make_item(admission, description="com percurso export")
        ws = _workbook(reader_client.get(EXPORT_URL))
        assert _rows(ws)[0][4] == "PRONTO SOCORRO"


@pytest.mark.django_db
class TestExportEnvelope:
    def test_mime_disposition_filename_and_cache(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_ENV")
        _make_item(admission, description="envelope export")
        response = reader_client.get(EXPORT_URL)
        assert response.status_code == 200
        assert response["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert response["Content-Disposition"] == (
            'attachment; filename="pendencias-20260610-0800.xlsx"'
        )
        _assert_private_no_store(response)

    def test_empty_result_is_valid_and_protected(
        self, reader_client: Client
    ) -> None:
        response = reader_client.get(f"{EXPORT_URL}?unidade=INEXISTENTE")
        ws = _workbook(response)
        assert ws is not None
        assert [cell.value for cell in ws[1]] == EXPECTED_HEADERS
        assert ws.max_row == 1
        _assert_private_no_store(response)

    def test_no_census_exports_headers_only(
        self, reader_client: Client
    ) -> None:
        patient = _make_patient("P_EXP_NOCENSUS")
        admission = _make_admission(patient, "ADM_EXP_NOCENSUS")
        _make_item(admission, description="sem foto export")
        ws = _workbook(reader_client.get(EXPORT_URL))
        assert ws is not None
        assert ws.max_row == 1
        assert _descriptions(ws) == []

    def test_formula_like_text_stays_literal(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_FORM")
        payload = "=HIPERLINK(\"http://x.example\";\"x\")"
        _make_item(admission, description=payload)
        other = _make_item(admission, description="+1+1")
        assert other.pk != 0
        ws = _workbook(reader_client.get(EXPORT_URL))
        assert ws is not None
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                assert cell.data_type in ("s", "inlineStr")
        assert payload in _descriptions(ws)

    def test_gets_never_mutate_business_tables(
        self, reader_client: Client
    ) -> None:
        admission = _eligible_admission("P_EXP_PURE")
        item = _make_item(admission, description="pura export")
        before_items = PendingItem.objects.count()
        before_revisions = PendingItemRevision.objects.count()
        reader_client.get(EXPORT_URL)
        reader_client.get(f"{EXPORT_URL}?view=items&status=all")
        reader_client.get(f"{EXPORT_URL}?min_dias=-1")
        assert PendingItem.objects.count() == before_items
        assert PendingItemRevision.objects.count() == before_revisions
        item.refresh_from_db()
        assert item.description == "pura export"
        assert item.version == 1
